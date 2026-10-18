import re
import os
from tkinter import Tk
from tkinter.filedialog import askopenfilename

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# =====================================================
# PILIH FILE
# =====================================================
root = Tk()
root.withdraw()
root.attributes('-topmost', True)

txt_file = askopenfilename(
    title="Pilih File MML TXT",
    filetypes=[("Text Files", "*.txt"), ("All Files", "*.*")]
)
root.destroy()

if not txt_file:
    print("Tidak ada file dipilih")
    exit()

base_name = os.path.splitext(os.path.basename(txt_file))[0]
output_file = os.path.join(os.path.dirname(txt_file), f"{base_name}.xlsx")

print(f"Input : {txt_file}")
print(f"Output : {output_file}")

# =====================================================
# BACA FILE
# =====================================================
with open(txt_file, "r", encoding="utf-8", errors="ignore") as f:
    lines = f.readlines()

# =====================================================
# CREATE WORKBOOK - FIX Pylance
# =====================================================
wb = Workbook()
# FIX Error 1: wb.active bisa None
if wb.active is not None:
    wb.remove(wb.active)

summary_ws = wb.create_sheet("SUMMARY")
sheet_dict = {}
summary_count = {}

RE_CMD = re.compile(r'^[A-Z0-9_ ]+:\s*;$')
RE_SEP = re.compile(r'^-{10,}\s*$')
RE_RESULT = re.compile(r'^\(Number of results')

i = 0
while i < len(lines):
    line = lines[i].strip()
    if RE_CMD.match(line):
        neinfo = line
        command_name = re.sub(r'[:;\s]+', '', line).strip()
        sheet_name = re.sub(r'[:\\/?*\[\]]', '_', command_name)[:31]
        if not sheet_name:
            i += 1
            continue

        nename = ""
        for x in range(i + 1, min(i + 10, len(lines))):
            temp = lines[x].strip()
            if "#MC" in temp and not temp.startswith("+++"):
                nename = temp
                break

        j = i + 1
        sep_index = -1
        while j < len(lines):
            if RE_SEP.match(lines[j]):
                sep_index = j
                break
            if RE_CMD.match(lines[j].strip()):
                break
            j += 1

        if sep_index == -1:
            i += 1
            continue

        table_type = ""
        for t in range(sep_index - 1, i, -1):
            if lines[t].strip():
                table_type = lines[t].strip()
                break

        header_line = lines[sep_index + 1].strip()
        if not header_line:
            i = sep_index + 1
            continue

        headers = re.split(r"\s{2,}", header_line)
        headers = ["NEName", "NEInfo", "TableType"] + headers

        if sheet_name not in sheet_dict:
            ws = wb.create_sheet(sheet_name)
            sheet_dict[sheet_name] = ws
            for col_num, h in enumerate(headers, 1):
                ws.cell(row=1, column=col_num, value=h).font = Font(bold=True)
        else:
            ws = sheet_dict[sheet_name]

        # FIX Error 2: max_row bisa None
        start_row = (ws.max_row or 0) + 1
        k = sep_index + 2
        data_count = 0
        while k < len(lines):
            row_strip = lines[k].strip()
            if not row_strip:
                k += 1
                continue
            if RE_RESULT.match(row_strip):
                k += 1
                continue
            if RE_SEP.match(lines[k]) or row_strip.startswith("---"):
                break
            if RE_CMD.match(row_strip):
                break

            cols = re.split(r"\s{2,}", row_strip)
            row_data = [nename, neinfo, table_type] + cols
            for col_num, v in enumerate(row_data, 1):
                ws.cell(row=start_row, column=col_num, value=v)
            start_row += 1
            data_count += 1
            k += 1

        summary_count[sheet_name] = summary_count.get(sheet_name, 0) + data_count
        i = k
    else:
        i += 1

summary_ws.append(["Sheet Name", "Rows"])
for cell in summary_ws[1]:
    cell.font = Font(bold=True)

for s_name, rows in sorted(summary_count.items()):
    r = (summary_ws.max_row or 0) + 1
    safe = s_name.replace("'", "''")
    c = summary_ws.cell(row=r, column=1, value=s_name)
    c.hyperlink = f"#'{safe}'!A1"
    c.style = "Hyperlink"
    summary_ws.cell(row=r, column=2, value=rows)

for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    for col_cells in ws.columns:
        max_len = 0
        # FIX: column bisa None
        col_idx = col_cells[0].column or 1
        letter = get_column_letter(col_idx)
        for cell in col_cells[:200]:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max_len + 2, 50)

wb.save(output_file)
print(f"\nExcel berhasil dibuat : {output_file}")
print(f"Total Sheet : {len(sheet_dict)}")
